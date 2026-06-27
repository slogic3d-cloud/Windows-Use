"""ZERO 2.0 — MCP-gebundene Tools für Gmail, Wix und Microsoft Office.

Einrichtung:
  Gmail:  GMAIL_ACCESS_TOKEN env-Variable oder Agent(secrets={"gmail_token": "..."})
  Wix:    WIX_API_KEY + WIX_ACCOUNT_ID env-Variablen oder Agent(secrets={...})
  Office: python-docx + openpyxl installieren: pip install python-docx openpyxl

Dateien werden in ./zero_downloads/ gespeichert und via /download/<datei> bereitgestellt.
"""

import json
import os
from datetime import datetime
from pathlib import Path
from typing import Literal

import httpx
from pydantic import Field

from windows_use.agent.tools.views import SharedBaseModel
from windows_use.tools import Tool

DOWNLOADS_DIR = Path("zero_downloads")
DOWNLOADS_DIR.mkdir(exist_ok=True)


# ─── Pydantic-Modelle ────────────────────────────────────────────────────────


class GmailRead(SharedBaseModel):
    max_results: int = Field(default=10, description="Maximale Anzahl ungelesener Mails (1–50)")
    query: str = Field(
        default="is:unread",
        description="Gmail-Suchquery, z.B. 'is:unread from:kunde@example.com'",
    )


class GmailSend(SharedBaseModel):
    to: str = Field(..., description="Empfänger-E-Mail-Adresse")
    subject: str = Field(..., description="Betreff der E-Mail")
    body: str = Field(..., description="E-Mail-Text (Plaintext oder HTML)")
    html: bool = Field(default=False, description="True wenn body HTML enthält")


class WixOrders(SharedBaseModel):
    status: Literal["PENDING", "APPROVED", "CANCELLED", "ALL"] = Field(
        default="ALL",
        description="Bestellstatus-Filter",
    )
    limit: int = Field(default=20, description="Maximale Anzahl Bestellungen (1–100)")


class WixProduct(SharedBaseModel):
    product_id: str | None = Field(
        default=None,
        description="Produkt-ID (leer = alle Produkte auflisten)",
    )
    action: Literal["get", "list", "update_stock"] = Field(
        default="list",
        description="Aktion: 'get' (einzelnes Produkt), 'list' (alle), 'update_stock' (Lager)",
    )
    new_stock: int | None = Field(
        default=None,
        description="Neuer Lagerbestand (nur für update_stock)",
    )


class OfficeCreate(SharedBaseModel):
    file_type: Literal["word", "excel"] = Field(
        ..., description="Dokumenttyp: 'word' (.docx) oder 'excel' (.xlsx)"
    )
    filename: str = Field(..., description="Dateiname ohne Erweiterung, z.B. 'Angebot_Kunde_A'")
    content: str = Field(
        ...,
        description=(
            "Für Word: Markdown-Text der in das Dokument geschrieben wird. "
            "Für Excel: JSON-Array von Zeilen, z.B. [[\"Name\",\"Wert\"],[\"A\",1]]"
        ),
    )
    sheet_name: str = Field(default="Tabelle1", description="Excel-Tabellenblatt-Name")


class OfficeEdit(SharedBaseModel):
    filename: str = Field(
        ..., description="Dateiname in zero_downloads/ (mit Erweiterung)"
    )
    instruction: str = Field(
        ...,
        description="Was geändert werden soll, z.B. 'Ersetze Firma X durch Firma Y'",
    )


# ─── Hilfsfunktionen ─────────────────────────────────────────────────────────


def _gmail_token(**kwargs) -> str:
    secrets: dict = kwargs.get("secrets", {})
    token = secrets.get("gmail_token") or os.environ.get("GMAIL_ACCESS_TOKEN", "")
    if not token:
        raise ValueError(
            "Kein Gmail-Token gefunden. Setze GMAIL_ACCESS_TOKEN oder "
            "übergib secrets={'gmail_token': '...'} an den Agent."
        )
    return token


def _wix_headers(**kwargs) -> dict:
    secrets: dict = kwargs.get("secrets", {})
    api_key = secrets.get("wix_api_key") or os.environ.get("WIX_API_KEY", "")
    account_id = secrets.get("wix_account_id") or os.environ.get("WIX_ACCOUNT_ID", "")
    if not api_key:
        raise ValueError(
            "Kein Wix-API-Key. Setze WIX_API_KEY oder übergib secrets={'wix_api_key': '...'}"
        )
    return {
        "Authorization": api_key,
        "wix-account-id": account_id,
        "Content-Type": "application/json",
    }


def _ts_filename(base: str, ext: str) -> str:
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"{ts}_{base}.{ext}"


# ─── Gmail Tools ─────────────────────────────────────────────────────────────


@Tool("gmail_read_tool", model=GmailRead)
async def gmail_read_tool(max_results: int = 10, query: str = "is:unread", **kwargs) -> str:
    """Liest E-Mails aus Gmail. Gibt Betreff, Absender und Vorschau der neuesten Mails zurück."""
    token = _gmail_token(**kwargs)
    headers = {"Authorization": f"Bearer {token}"}
    base = "https://gmail.googleapis.com/gmail/v1/users/me"

    async with httpx.AsyncClient(timeout=30) as client:
        r = await client.get(
            f"{base}/messages",
            headers=headers,
            params={"q": query, "maxResults": max_results},
        )
        r.raise_for_status()
        msg_ids = [m["id"] for m in r.json().get("messages", [])]

        results = []
        for mid in msg_ids[:max_results]:
            mr = await client.get(
                f"{base}/messages/{mid}",
                headers=headers,
                params={"format": "metadata", "metadataHeaders": ["From", "Subject", "Date"]},
            )
            mr.raise_for_status()
            payload = mr.json()
            headers_list = payload.get("payload", {}).get("headers", [])
            meta = {h["name"]: h["value"] for h in headers_list}
            snippet = payload.get("snippet", "")[:200]
            results.append(
                f"**Von:** {meta.get('From', '?')}\n"
                f"**Betreff:** {meta.get('Subject', '?')}\n"
                f"**Datum:** {meta.get('Date', '?')}\n"
                f"**Vorschau:** {snippet}"
            )

    if not results:
        return "Keine E-Mails gefunden."
    return f"**{len(results)} E-Mail(s) gefunden:**\n\n" + "\n\n---\n\n".join(results)


@Tool("gmail_send_tool", model=GmailSend)
async def gmail_send_tool(to: str, subject: str, body: str, html: bool = False, **kwargs) -> str:
    """Sendet eine E-Mail über Gmail. Erfordert Bestätigung (ROT-Tier)."""
    import base64
    from email.mime.multipart import MIMEMultipart
    from email.mime.text import MIMEText

    token = _gmail_token(**kwargs)
    msg = MIMEMultipart("alternative") if html else MIMEText(body, "plain", "utf-8")
    msg["To"] = to
    msg["Subject"] = subject
    if html:
        msg.attach(MIMEText(body, "html", "utf-8"))

    raw = base64.urlsafe_b64encode(msg.as_bytes()).decode()
    headers = {
        "Authorization": f"Bearer {token}",
        "Content-Type": "application/json",
    }
    async with httpx.AsyncClient(timeout=30) as client:
        r = await client.post(
            "https://gmail.googleapis.com/gmail/v1/users/me/messages/send",
            headers=headers,
            json={"raw": raw},
        )
        r.raise_for_status()

    return f"E-Mail erfolgreich an '{to}' gesendet. Betreff: '{subject}'"


# ─── Wix Tools ───────────────────────────────────────────────────────────────


@Tool("wix_orders_tool", model=WixOrders)
async def wix_orders_tool(
    status: str = "ALL", limit: int = 20, **kwargs
) -> str:
    """Ruft Bestellungen aus dem S-LOGIC 3D Wix-Shop ab."""
    headers = _wix_headers(**kwargs)
    site_id = kwargs.get("secrets", {}).get("wix_site_id") or os.environ.get("WIX_SITE_ID", "")
    filter_body: dict = {"limit": limit, "offset": 0}
    if status != "ALL":
        filter_body["filter"] = {"status": status}

    async with httpx.AsyncClient(timeout=30) as client:
        r = await client.post(
            f"https://www.wixapis.com/ecom/v1/orders/search",
            headers={**headers, "wix-site-id": site_id},
            json={"search": filter_body},
        )
        r.raise_for_status()
        orders = r.json().get("orders", [])

    if not orders:
        return "Keine Bestellungen gefunden."

    lines = [f"**{len(orders)} Bestellung(en):**\n"]
    for o in orders:
        oid = o.get("id", "?")[:8]
        buyer = o.get("buyerInfo", {}).get("email", "?")
        total = o.get("priceSummary", {}).get("total", {}).get("formattedAmount", "?")
        st = o.get("status", "?")
        created = o.get("createdDate", "?")[:10]
        lines.append(f"- `{oid}` | {buyer} | {total} | {st} | {created}")

    return "\n".join(lines)


@Tool("wix_product_tool", model=WixProduct)
async def wix_product_tool(
    product_id: str | None = None,
    action: str = "list",
    new_stock: int | None = None,
    **kwargs,
) -> str:
    """Ruft Produkte aus dem Wix-Shop ab oder aktualisiert den Lagerbestand."""
    headers = _wix_headers(**kwargs)
    site_id = kwargs.get("secrets", {}).get("wix_site_id") or os.environ.get("WIX_SITE_ID", "")
    h = {**headers, "wix-site-id": site_id}

    async with httpx.AsyncClient(timeout=30) as client:
        if action == "list":
            r = await client.post(
                "https://www.wixapis.com/stores/v1/products/query",
                headers=h,
                json={"query": {"paging": {"limit": 50}}},
            )
            r.raise_for_status()
            products = r.json().get("products", [])
            lines = [f"**{len(products)} Produkt(e):**\n"]
            for p in products:
                lines.append(
                    f"- `{p.get('id','?')[:8]}` | {p.get('name','?')} | "
                    f"{p.get('priceData',{}).get('formattedPrice','?')} | "
                    f"Lager: {p.get('stock',{}).get('quantity','?')}"
                )
            return "\n".join(lines)

        elif action == "get" and product_id:
            r = await client.get(
                f"https://www.wixapis.com/stores/v1/products/{product_id}",
                headers=h,
            )
            r.raise_for_status()
            p = r.json().get("product", {})
            return json.dumps(p, ensure_ascii=False, indent=2)

        elif action == "update_stock" and product_id and new_stock is not None:
            r = await client.post(
                "https://www.wixapis.com/stores/v1/inventoryItems/updateInventoryVariants",
                headers=h,
                json={
                    "inventoryItem": {
                        "externalId": product_id,
                        "variants": [{"quantity": new_stock}],
                    }
                },
            )
            r.raise_for_status()
            return f"Lagerbestand für Produkt '{product_id}' auf {new_stock} gesetzt."

        return "Ungültige Kombination aus action/product_id/new_stock."


# ─── Office Tools ────────────────────────────────────────────────────────────


@Tool("office_create_tool", model=OfficeCreate)
async def office_create_tool(
    file_type: str, filename: str, content: str, sheet_name: str = "Tabelle1", **kwargs
) -> str:
    """Erstellt ein Word- (.docx) oder Excel-Dokument (.xlsx) und stellt es zum Download bereit.

    Das Dokument wird in zero_downloads/ gespeichert.
    Rückgabe enthält den Download-Pfad.
    """
    safe_name = "".join(c for c in filename if c.isalnum() or c in "_-. ")
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")

    if file_type == "word":
        try:
            from docx import Document
            from docx.shared import Pt
        except ImportError:
            return "python-docx nicht installiert. Bitte: pip install python-docx"

        doc = Document()
        doc.core_properties.author = "ZERO 2.0"
        for line in content.split("\n"):
            line = line.strip()
            if line.startswith("# "):
                doc.add_heading(line[2:], level=1)
            elif line.startswith("## "):
                doc.add_heading(line[3:], level=2)
            elif line.startswith("### "):
                doc.add_heading(line[4:], level=3)
            elif line.startswith("- ") or line.startswith("* "):
                doc.add_paragraph(line[2:], style="List Bullet")
            elif line:
                para = doc.add_paragraph()
                # Fettdruck für **text**
                parts = line.split("**")
                for i, part in enumerate(parts):
                    run = para.add_run(part)
                    run.bold = (i % 2 == 1)
                    run.font.size = Pt(11)

        out_name = f"{ts}_{safe_name}.docx"
        out_path = DOWNLOADS_DIR / out_name
        doc.save(str(out_path))
        return f"Word-Dokument erstellt: /download/{out_name}"

    elif file_type == "excel":
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font
        except ImportError:
            return "openpyxl nicht installiert. Bitte: pip install openpyxl"

        try:
            rows = json.loads(content)
            if not isinstance(rows, list):
                raise ValueError
        except (json.JSONDecodeError, ValueError):
            return (
                "Excel-Content muss ein JSON-Array sein, z.B. "
                '[[\"Name\",\"Wert\"],[\"A\",1]]'
            )

        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name
        for r_idx, row in enumerate(rows, 1):
            for c_idx, val in enumerate(row, 1):
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                if r_idx == 1:
                    cell.font = Font(bold=True)

        out_name = f"{ts}_{safe_name}.xlsx"
        out_path = DOWNLOADS_DIR / out_name
        wb.save(str(out_path))
        return f"Excel-Tabelle erstellt: /download/{out_name}"

    return f"Unbekannter Dateityp: {file_type}"


@Tool("office_edit_tool", model=OfficeEdit)
async def office_edit_tool(filename: str, instruction: str, **kwargs) -> str:
    """Bearbeitet ein vorhandenes Word-Dokument (Textersetzung). Erfordert Bestätigung (ROT-Tier)."""
    try:
        from docx import Document
    except ImportError:
        return "python-docx nicht installiert. Bitte: pip install python-docx"

    path = DOWNLOADS_DIR / filename
    if not path.exists():
        return f"Datei nicht gefunden: zero_downloads/{filename}"

    doc = Document(str(path))

    # Einfache Ersetzungslogik: "Ersetze X durch Y"
    instruction_lower = instruction.lower()
    if "ersetze" in instruction_lower and " durch " in instruction_lower:
        parts = instruction_lower.split("ersetze ", 1)[1].split(" durch ", 1)
        if len(parts) == 2:
            old_text = parts[0].strip().strip("\"'")
            new_text = parts[1].strip().strip("\"'")
            count = 0
            for para in doc.paragraphs:
                if old_text in para.text:
                    for run in para.runs:
                        if old_text in run.text:
                            run.text = run.text.replace(old_text, new_text)
                            count += 1
            doc.save(str(path))
            return f"'{old_text}' wurde {count}x durch '{new_text}' ersetzt in {filename}."

    return (
        f"Anweisung konnte nicht automatisch ausgeführt werden: '{instruction}'. "
        "Bitte in der Form 'Ersetze X durch Y' formulieren."
    )


# ─── Tool-Liste für Agent-Registrierung ──────────────────────────────────────

MCP_TOOLS = [
    gmail_read_tool,
    gmail_send_tool,
    wix_orders_tool,
    wix_product_tool,
    office_create_tool,
    office_edit_tool,
]
